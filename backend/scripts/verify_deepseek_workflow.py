"""真实 DeepSeek 验收：文字、图片、知识库、生成、修改和文件下载（调用计费 API）。"""

import asyncio
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import cv2
import httpx
import numpy as np
from docx import Document
from fastapi import HTTPException
from pptx import Presentation

from app.core.llm.factory import get_llm

BASE = "http://127.0.0.1:5173"
EVIDENCE = Path(__file__).resolve().parents[2] / ".local" / "deepseek-verification.json"


async def main():
    checks = {}
    image_path = EVIDENCE.parent / "deepseek-red-test.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image = np.zeros((64, 64, 3), dtype=np.uint8)
    image[:, :, 2] = 255
    success, encoded = cv2.imencode(".png", image)
    assert success
    image_path.write_bytes(encoded.tobytes())
    description = await get_llm().vision(
        str(image_path), "请用中文说明图片的主要颜色，只回复颜色名称。"
    )
    assert "红" in description, description
    checks["vision"] = description
    print("PASS: DeepSeek image understanding", flush=True)

    async with httpx.AsyncClient(base_url=BASE, timeout=240, trust_env=False) as client:

        async def post(path, payload):
            response = await client.post(path, json=payload)
            if response.status_code != 200:
                raise RuntimeError(f"{path}: HTTP {response.status_code}; {response.text[:300]}")
            return response.json()

        doc = Document()
        doc.add_heading("同分母分数加法讲义", 0)
        doc.add_paragraph("同分母分数相加，分母不变，分子相加。例：1/5 + 2/5 = 3/5。")
        doc.add_paragraph("用等分圆形帮助小学五年级学生理解：每份大小相同，合并份数。")
        buffer = io.BytesIO()
        doc.save(buffer)
        upload = await client.post(
            "/api/v1/documents/upload",
            files={
                "file": (
                    "DeepSeek验收_同分母分数.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            },
            data={"is_reference": "true"},
        )
        upload.raise_for_status()
        doc_id = upload.json()["id"]
        for _ in range(45):
            response = await client.get("/api/v1/documents")
            state = next(
                item["status"] for item in response.json()["documents"] if item["id"] == doc_id
            )
            if state != "处理中":
                break
            await asyncio.sleep(1)
        assert state in {"已完成", "有冲突"}, state
        hits = await post("/api/v1/knowledge/search", {"query": "同分母分数加法", "k": 5})
        assert any(hit["doc_id"] == doc_id for hit in hits["hits"])
        checks["knowledge"] = {"document_id": doc_id, "status": state}
        print("PASS: real upload, extraction and local retrieval", flush=True)

        chat = await post(
            "/api/v1/chat",
            {
                "messages": [
                    {
                        "role": "user",
                        "content": "给小学五年级讲同分母分数加法，20分钟，简约风格。目标：理解分母不变分子相加；"
                        "重点是用等分图理解加法。请生成4页课件和简短教案，不需要互动游戏，开始生成。",
                    }
                ],
                "granularity": "快速",
                "reference_doc_ids": [doc_id],
            },
        )
        artifacts = chat["artifacts"]
        assert artifacts and artifacts["ppt"]["slides"] and artifacts["word"]["data"]
        assert "分数" in artifacts["intent"]["topic"]
        assert artifacts["knowledge_hits"]
        for kind in ("ppt", "word"):
            response = await client.get(f"/api/v1/files/{artifacts[kind]['filename']}")
            response.raise_for_status()
            if kind == "ppt":
                assert len(Presentation(io.BytesIO(response.content)).slides) > 0
            else:
                assert len(Document(io.BytesIO(response.content)).paragraphs) > 0
        checks["chat"] = {
            "topic": artifacts["intent"]["topic"],
            "slides": len(artifacts["ppt"]["slides"]),
            "downloads": "PASS",
        }
        print("PASS: real lesson, PPT/Word downloads and references", flush=True)

        ppt = await post(
            "/api/v1/revise",
            {"slides": artifacts["ppt"]["slides"], "feedback": "增加一道1/5+2/5的课堂练习"},
        )
        word = await post(
            "/api/v1/revise/word",
            {"word": artifacts["word"]["data"], "feedback": "作业增加一道2/7+3/7的计算题"},
        )
        assert ppt["slides"] and word["word"]
        checks["revise"] = "PASS"
        print("PASS: real PPT and Word revisions", flush=True)

        exam = await post("/api/v1/exam/generate", {"intent": artifacts["intent"], "n": 3})
        assert len(exam["questions"]) >= 3 and exam["bank_saved"] >= 3
        exam_download = await client.get(f"/api/v1/files/{exam['filename']}")
        exam_download.raise_for_status()
        assert Document(io.BytesIO(exam_download.content)).paragraphs
        checks["exam"] = {"questions": len(exam["questions"]), "download": "PASS"}
        print("PASS: real exam, bank and download", flush=True)

        interactive = await post("/api/v1/interactive/generate", {"intent": artifacts["intent"]})
        assert interactive["filename"].endswith(".html")
        checks["interactive"] = "PASS"
        print("PASS: real interactive HTML generation", flush=True)
    EVIDENCE.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Live verification saved to {EVIDENCE}", flush=True)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except HTTPException as error:
        print(f"FAIL: {error.detail}")
        raise SystemExit(1) from None
