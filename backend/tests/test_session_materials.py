import io
import uuid
from pathlib import Path

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.db import init_db
from app.main import app

client = TestClient(app)
init_db()


@pytest.fixture(autouse=True)
def isolated_files(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)


def headers(session):
    return {"X-EduMind-Session": session}


def upload(session, name):
    document = WordDocument()
    document.add_paragraph("TCP 三次握手 SYN ACK，滑动窗口与四次挥手。")
    content = io.BytesIO()
    document.save(content)
    response = client.post(
        "/api/v1/documents/upload",
        headers=headers(session),
        files={"file": (name, content.getvalue())},
    )
    assert response.status_code == 200
    return response.json()["id"]


def test_documents_graph_retrieval_and_deletion_are_scoped():
    a, b = "a_" + uuid.uuid4().hex, "b_" + uuid.uuid4().hex
    doc_a = upload(a, "A讲义.docx")
    doc_b = upload(b, "B讲义.docx")
    listed = client.get("/api/v1/documents", headers=headers(a)).json()["documents"]
    assert {doc["id"] for doc in listed} == {doc_a}
    assert client.get("/api/v1/knowledge/graph", headers=headers(a)).json()["nodes"]
    hits = client.post(
        "/api/v1/knowledge/search", headers=headers(a), json={"query": "TCP", "k": 5}
    ).json()["hits"]
    assert hits and {hit["doc_id"] for hit in hits} == {doc_a}
    denied = client.delete(f"/api/v1/documents/{doc_a}", headers=headers(b))
    assert denied.status_code == 404
    removed = client.delete(f"/api/v1/documents/{doc_a}", headers=headers(a))
    assert removed.status_code == 200
    assert client.get("/api/v1/documents", headers=headers(a)).json()["documents"] == []
    assert (
        client.post("/api/v1/knowledge/search", headers=headers(a), json={"query": "TCP"}).json()[
            "hits"
        ]
        == []
    )
    assert {
        doc["id"] for doc in client.get("/api/v1/documents", headers=headers(b)).json()["documents"]
    } == {doc_b}


def test_generated_files_belong_to_their_session_and_session_delete_cleans():
    session = "files_" + uuid.uuid4().hex
    response = client.post(
        "/api/v1/chat",
        headers=headers(session),
        json={
            "messages": [{"role": "user", "content": "讲 TCP，开始生成"}],
        },
    )
    assert response.status_code == 200
    filename = response.json()["artifacts"]["ppt"]["filename"]
    assert client.get(f"/api/v1/files/{filename}?session_id={session}").status_code == 200
    assert client.get(f"/api/v1/files/{filename}?session_id=other").status_code == 404
    assert client.delete(f"/api/v1/sessions/{session}", headers=headers("other")).status_code == 403
    assert client.delete(f"/api/v1/sessions/{session}", headers=headers(session)).status_code == 200
    assert not (Path("data/output") / filename).exists()


def test_busy_document_cannot_be_deleted_and_invalid_scope_rejected():
    from app.db import SessionLocal
    from app.db.models import Document

    session = "busy_" + uuid.uuid4().hex
    with SessionLocal() as db:
        doc = Document(
            session_id=session,
            user_id="default",
            filename="解析中的教材",
            file_path="unused",
            file_type="docx",
        )
        db.add(doc)
        db.commit()
        doc_id = doc.id
    assert client.delete(f"/api/v1/documents/{doc_id}", headers=headers(session)).status_code == 409
    assert client.delete(f"/api/v1/sessions/{session}", headers=headers(session)).status_code == 409
    assert client.get("/api/v1/documents", headers=headers("../bad")).status_code == 400
    assert (
        client.get("/api/v1/documents?session_id=other", headers=headers(session)).status_code
        == 400
    )


def test_conflict_review_and_graph_expansion_are_scoped():
    from app.core.session_context import session_scope
    from app.db import SessionLocal
    from app.db.models import Conflict, KnowledgeEdge, KnowledgeNode
    from app.knowledge.graph import neighborhood

    a, b = "graph_" + uuid.uuid4().hex, "other_" + uuid.uuid4().hex
    with SessionLocal() as db:
        first = KnowledgeNode(session_id=a, user_id="default", title="当前", content="当前")
        foreign = KnowledgeNode(session_id=b, user_id="default", title="其他", content="其他")
        db.add_all([first, foreign])
        db.flush()
        conflict = Conflict(
            session_id=b,
            user_id="default",
            existing_knowledge={"id": foreign.id},
            new_knowledge={"title": "新知识"},
        )
        db.add(conflict)
        db.add(
            KnowledgeEdge(
                session_id=b,
                user_id="default",
                from_node=first.id,
                to_node=foreign.id,
                relation_type="相关关联",
            )
        )
        db.commit()
        first_id, conflict_id = first.id, conflict.id
    assert client.get("/api/v1/conflicts", headers=headers(a)).json()["conflicts"] == []
    assert (
        client.post(
            f"/api/v1/conflicts/{conflict_id}/review", headers=headers(a), json={"action": "保留旧"}
        ).status_code
        == 404
    )
    token = session_scope.set(a)
    try:
        assert neighborhood([first_id]) == []
    finally:
        session_scope.reset(token)


def test_duplicate_sources_survive_deleting_one_document():
    session = "shared_" + uuid.uuid4().hex
    first = upload(session, "第一份.docx")
    second = upload(session, "第二份.docx")
    assert client.delete(f"/api/v1/documents/{first}", headers=headers(session)).status_code == 200
    assert client.get("/api/v1/knowledge/graph", headers=headers(session)).json()["nodes"]
    assert client.delete(f"/api/v1/documents/{second}", headers=headers(session)).status_code == 200
    assert client.get("/api/v1/knowledge/graph", headers=headers(session)).json()["nodes"] == []


def test_legacy_files_import_is_idempotent_and_cannot_reassign():
    folder = Path("data/output")
    folder.mkdir(parents=True)
    (folder / "old.docx").write_bytes(b"old content")
    a, b = "old_" + uuid.uuid4().hex, "new_" + uuid.uuid4().hex
    for scope in (a, a, b):
        assert (
            client.post(
                "/api/v1/sessions/import-files",
                headers=headers(scope),
                json={"filenames": ["old.docx"]},
            ).status_code
            == 200
        )
    assert client.get(f"/api/v1/files/old.docx?session_id={a}").status_code == 200
    assert client.get(f"/api/v1/files/old.docx?session_id={b}").status_code == 404
